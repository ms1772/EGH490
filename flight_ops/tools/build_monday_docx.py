#!/usr/bin/env python3
"""Render flight_ops/runsheets/monday_steps.md as the flight-day Word document.

  python flight_ops/tools/build_monday_docx.py

Writes MONDAY_O134_Step_By_Step.docx into the project folder (the parent of the
repo), never into the repo: the repository carries code, not documents.

CLOSE THE DOCX IN WORD FIRST. python-docx cannot overwrite a file Word holds
open, and the old builder quietly diverted to ..._UPDATED.docx instead of
saying so -- which is how you end up carrying yesterday's run-sheet into the
lab while a newer one sits beside it.
"""
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
PROJECT = REPO.parent

_spec = importlib.util.spec_from_file_location("md_to_docx", HERE / "md_to_docx.py")
md_to_docx = importlib.util.module_from_spec(_spec)
sys.modules["md_to_docx"] = md_to_docx
_spec.loader.exec_module(md_to_docx)


def main() -> int:
    src = REPO / "flight_ops" / "runsheets" / "monday_steps.md"
    if not src.exists():
        print(f"missing {src}", file=sys.stderr)
        return 1

    out_name = "MONDAY_O134_Step_By_Step.docx"
    target = PROJECT / out_name
    if target.exists():
        try:
            with target.open("r+b"):
                pass
        except OSError:
            print(f"{target} is open in Word -- close it and run this again.",
                  file=sys.stderr)
            return 1

    out = md_to_docx.build(
        out_name,
        "O-134 Flight Session",
        "Step by step, from motion capture to the DECK-GA and QuickNav flights",
        "Follow in order. Stop on any step that fails twice.",
        src,
        [],
    )
    print(f"wrote {out} ({out.stat().st_size // 1024} KB)")

    from docx import Document
    d = Document(out)
    code = sum(1 for p in d.paragraphs
               if p.runs and p.runs[0].font.name == "Consolas")
    print(f"paragraphs: {len(d.paragraphs)}  code blocks: {code}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
